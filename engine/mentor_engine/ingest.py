"""Text extraction and chunking for knowledge base, projects and indexed folders."""

from __future__ import annotations

import csv
import io
import re
from pathlib import Path

TEXT_EXT = {
    ".txt", ".md", ".markdown", ".rst", ".csv", ".tsv", ".json", ".yaml", ".yml", ".xml", ".html", ".htm",
    ".py", ".js", ".ts", ".tsx", ".jsx", ".java", ".cs", ".go", ".rb", ".php", ".sql", ".sh", ".ps1",
    ".css", ".ini", ".toml", ".cfg", ".log",
}
DOC_EXT = {".pdf", ".docx", ".xlsx", ".pptx"}
SUPPORTED = TEXT_EXT | DOC_EXT
SKIP_DIRS = {
    "node_modules", "__pycache__", ".git", ".svn", ".hg", "venv", ".venv", "env", "build", "dist", "out",
    ".cache", ".idea", ".vscode", "target", ".next", ".pytest_cache", ".mypy_cache",
}
MAX_FILE_BYTES = 25 * 1024 * 1024


def skip_reason(path: Path, root: Path | None = None) -> str | None:
    rel = path.relative_to(root) if root else path
    if any(part in SKIP_DIRS for part in rel.parts):
        return "system folder"
    if path.name.startswith("~$") or path.name.startswith("."):
        return "system file"
    if path.suffix.lower() not in SUPPORTED:
        return "unsupported type"
    try:
        if path.stat().st_size > MAX_FILE_BYTES:
            return "too large"
    except OSError:
        return "unreadable"
    return None


def extract_text(path: Path) -> str:
    ext = path.suffix.lower()
    if ext == ".pdf":
        from pypdf import PdfReader

        reader = PdfReader(str(path))
        return "\n\n".join(f"[page {i + 1}]\n{p.extract_text() or ''}" for i, p in enumerate(reader.pages))
    if ext == ".docx":
        import docx

        d = docx.Document(str(path))
        parts = [p.text for p in d.paragraphs if p.text.strip()]
        for table in d.tables:
            for row in table.rows:
                parts.append(" | ".join(c.text.strip() for c in row.cells))
        return "\n".join(parts)
    if ext == ".xlsx":
        import openpyxl

        wb = openpyxl.load_workbook(str(path), read_only=True, data_only=True)
        parts = []
        for ws in wb.worksheets:
            parts.append(f"[sheet {ws.title}]")
            for i, row in enumerate(ws.iter_rows(values_only=True)):
                if i > 5000:
                    parts.append("... (truncated)")
                    break
                if any(v is not None for v in row):
                    parts.append(" | ".join("" if v is None else str(v) for v in row))
        return "\n".join(parts)
    if ext == ".pptx":
        from pptx import Presentation

        prs = Presentation(str(path))
        parts = []
        for i, slide in enumerate(prs.slides, 1):
            parts.append(f"[slide {i}]")
            for shape in slide.shapes:
                if shape.has_text_frame:
                    parts.append(shape.text_frame.text)
        return "\n".join(parts)
    if ext in (".html", ".htm"):
        from bs4 import BeautifulSoup

        return BeautifulSoup(path.read_text("utf-8", errors="ignore"), "html.parser").get_text("\n")
    return path.read_text("utf-8", errors="ignore")


def extract_bytes(name: str, data: bytes) -> str:
    """Extract text from an in-memory upload by spilling to a temp file when needed."""
    import tempfile

    suffix = Path(name).suffix.lower()
    if suffix in TEXT_EXT:
        return data.decode("utf-8", errors="ignore")
    with tempfile.NamedTemporaryFile(suffix=suffix, delete=False) as f:
        f.write(data)
        tmp = Path(f.name)
    try:
        return extract_text(tmp)
    finally:
        tmp.unlink(missing_ok=True)


def chunk_text(text: str, size: int = 1200, overlap: int = 200) -> list[str]:
    """Paragraph-aware chunking with character overlap."""
    text = re.sub(r"\n{3,}", "\n\n", text).strip()
    if not text:
        return []
    paras = re.split(r"\n\s*\n", text)
    chunks: list[str] = []
    cur = ""
    for p in paras:
        if len(p) > size:
            for i in range(0, len(p), size - overlap):
                piece = p[i : i + size]
                if cur:
                    chunks.append(cur)
                    cur = ""
                chunks.append(piece)
            continue
        if len(cur) + len(p) + 2 > size and cur:
            chunks.append(cur)
            cur = cur[-overlap:] + "\n\n" + p if overlap else p
        else:
            cur = f"{cur}\n\n{p}" if cur else p
    if cur:
        chunks.append(cur)
    return [c.strip() for c in chunks if c.strip()]


def csv_preview(text: str, max_rows: int = 20) -> dict:
    rows = list(csv.reader(io.StringIO(text)))
    return {"header": rows[0] if rows else [], "rows": rows[1 : max_rows + 1], "total_rows": max(0, len(rows) - 1)}
