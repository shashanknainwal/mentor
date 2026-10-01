"""File generation (Word, Excel, CSV, Markdown, HTML, charts, images) and local file access."""

from __future__ import annotations

import asyncio
import csv
import html
import io
import re
import statistics
from pathlib import Path

from .. import artifacts, ingest
from ..knowledge import resolve_allowed
from . import ToolError, n, obj, s, tool

ARR_STR = {"type": "array", "items": {"type": "string"}}


def _saved(ctx, doc) -> str:
    ctx.created_artifacts.append(doc)
    return f"Saved artifact '{doc['name']}' ({doc['size']:,} bytes) at {doc['path']} [artifact:{doc['id']}]"


# ---------------------------------------------------------------------------
# Word
# ---------------------------------------------------------------------------
def markdown_to_docx(title: str, markdown: str) -> bytes:
    import docx
    from docx.shared import Pt, RGBColor

    d = docx.Document()
    style = d.styles["Normal"]
    style.font.name = "Arial"
    style.font.size = Pt(10.5)
    h = d.add_heading(title, level=0)
    for r in h.runs:
        r.font.color.rgb = RGBColor(0x00, 0x33, 0x8D)

    def add_runs(par, text):
        for part in re.split(r"(\*\*[^*]+\*\*|\*[^*]+\*|`[^`]+`)", text):
            if not part:
                continue
            if part.startswith("**") and part.endswith("**"):
                par.add_run(part[2:-2]).bold = True
            elif part.startswith("*") and part.endswith("*") and len(part) > 2:
                par.add_run(part[1:-1]).italic = True
            elif part.startswith("`") and part.endswith("`"):
                par.add_run(part[1:-1]).font.name = "Consolas"
            else:
                par.add_run(part)

    lines = markdown.splitlines()
    i = 0
    while i < len(lines):
        line = lines[i].rstrip()
        if line.startswith("|") and i + 1 < len(lines) and re.match(r"^\|?\s*:?-+", lines[i + 1].strip()):
            rows = [line]
            i += 2
            while i < len(lines) and lines[i].strip().startswith("|"):
                rows.append(lines[i])
                i += 1
            cells = [[c.strip() for c in r.strip().strip("|").split("|")] for r in rows]
            table = d.add_table(rows=len(cells), cols=max(len(r) for r in cells))
            table.style = "Light Grid Accent 1"
            for ri, row in enumerate(cells):
                for ci, val in enumerate(row):
                    table.cell(ri, ci).text = val
                    if ri == 0:
                        for r in table.cell(ri, ci).paragraphs[0].runs:
                            r.bold = True
            continue
        m = re.match(r"^(#{1,4})\s+(.*)", line)
        if m:
            d.add_heading(m.group(2), level=len(m.group(1)))
        elif re.match(r"^\s*[-*•]\s+", line):
            add_runs(d.add_paragraph(style="List Bullet"), re.sub(r"^\s*[-*•]\s+", "", line))
        elif re.match(r"^\s*\d+[.)]\s+", line):
            add_runs(d.add_paragraph(style="List Number"), re.sub(r"^\s*\d+[.)]\s+", "", line))
        elif line.strip():
            add_runs(d.add_paragraph(), line)
        i += 1
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


@tool(
    "create_docx",
    "Create a Word document (.docx) from markdown (headings, bullets, numbered lists, **bold**, tables). "
    "Use for reports, memos, proposals, letters.",
    obj(["title", "markdown"], title=s("Document title"), markdown=s("Body in markdown"), filename=s("Optional file name")),
    risk="medium",
    category="Documents",
)
async def create_docx(args, ctx):
    data = await asyncio.to_thread(markdown_to_docx, args["title"], args["markdown"])
    name = args.get("filename") or f"{args['title']}.docx"
    return _saved(ctx, artifacts.save(name if name.endswith(".docx") else name + ".docx", data, ctx.session_id, "create_docx"))


# ---------------------------------------------------------------------------
# Excel / CSV
# ---------------------------------------------------------------------------
@tool(
    "create_xlsx",
    "Create an Excel workbook. Each sheet has a name and rows (first row = header). Numbers stay numeric.",
    obj(
        ["filename", "sheets"],
        filename=s("File name, e.g. 'pipeline.xlsx'"),
        sheets={"type": "array", "items": obj(["name", "rows"], name=s("Sheet name"),
                                               rows={"type": "array", "items": {"type": "array", "items": {}}})},
    ),
    risk="medium",
    category="Documents",
)
async def create_xlsx(args, ctx):
    import openpyxl
    from openpyxl.styles import Font, PatternFill

    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for sh in args["sheets"]:
        ws = wb.create_sheet(str(sh["name"])[:31] or "Sheet")
        for r, row in enumerate(sh["rows"], 1):
            for c, val in enumerate(row, 1):
                if isinstance(val, str):
                    try:
                        val = float(val.replace(",", "")) if re.fullmatch(r"-?[\d,]+(\.\d+)?", val) else val
                    except ValueError:
                        pass
                cell = ws.cell(row=r, column=c, value=val)
                if r == 1:
                    cell.font = Font(bold=True, color="FFFFFF")
                    cell.fill = PatternFill("solid", fgColor="00338D")
        for col in ws.columns:
            ws.column_dimensions[col[0].column_letter].width = min(60, max(10, *(len(str(c.value or "")) + 2 for c in col)))
    buf = io.BytesIO()
    wb.save(buf)
    name = args["filename"] if args["filename"].endswith(".xlsx") else args["filename"] + ".xlsx"
    return _saved(ctx, artifacts.save(name, buf.getvalue(), ctx.session_id, "create_xlsx"))


@tool(
    "create_csv",
    "Create a CSV file from rows (first row = header).",
    obj(["filename", "rows"], filename=s("File name"), rows={"type": "array", "items": {"type": "array", "items": {}}}),
    risk="medium",
    category="Documents",
)
async def create_csv(args, ctx):
    buf = io.StringIO()
    csv.writer(buf).writerows(args["rows"])
    name = args["filename"] if args["filename"].endswith(".csv") else args["filename"] + ".csv"
    return _saved(ctx, artifacts.save(name, buf.getvalue(), ctx.session_id, "create_csv"))


@tool(
    "save_text_file",
    "Save text content as an artifact: markdown (.md), code (.py/.sql/.js...), JSON, Mermaid diagrams (.mmd), or plain text.",
    obj(["filename", "content"], filename=s("File name with extension"), content=s("Full file content")),
    risk="medium",
    category="Documents",
)
async def save_text_file(args, ctx):
    return _saved(ctx, artifacts.save(args["filename"], args["content"], ctx.session_id, "save_text_file"))


@tool(
    "create_html_report",
    "Create a styled, self-contained HTML report from markdown. Good for sharing in a browser.",
    obj(["title", "markdown"], title=s("Title"), markdown=s("Body in markdown")),
    risk="medium",
    category="Documents",
)
async def create_html_report(args, ctx):
    body = []
    for line in args["markdown"].splitlines():
        esc = html.escape(line)
        esc = re.sub(r"\*\*(.+?)\*\*", r"<strong>\1</strong>", esc)
        m = re.match(r"^(#{1,4})\s+(.*)", esc)
        if m:
            body.append(f"<h{len(m.group(1)) + 1}>{m.group(2)}</h{len(m.group(1)) + 1}>")
        elif re.match(r"^\s*[-*]\s+", esc):
            item = re.sub(r"^\s*[-*]\s+", "", esc)
            body.append(f"<li>{item}</li>")
        elif esc.strip():
            body.append(f"<p>{esc}</p>")
    doc = f"""<!doctype html><html><head><meta charset="utf-8"><title>{html.escape(args['title'])}</title>
<style>body{{font-family:Arial,Helvetica,sans-serif;max-width:860px;margin:40px auto;padding:0 20px;color:#1e1e1e;line-height:1.55}}
h1{{color:#00338D;border-bottom:3px solid #00338D;padding-bottom:8px}}h2,h3{{color:#00338D}}li{{margin:4px 0}}</style></head>
<body><h1>{html.escape(args['title'])}</h1>{''.join(body)}</body></html>"""
    return _saved(ctx, artifacts.save(f"{args['title']}.html", doc, ctx.session_id, "create_html_report"))


# ---------------------------------------------------------------------------
# Charts (dependency-free SVG)
# ---------------------------------------------------------------------------
PALETTE = ["#00338D", "#0091DA", "#483698", "#470A68", "#00A3A1", "#C6007E", "#6D2077", "#F68D2E"]


def svg_chart(kind: str, title: str, labels: list[str], series: list[dict]) -> str:
    W, H, P = 760, 420, 60
    out = [f'<svg xmlns="http://www.w3.org/2000/svg" width="{W}" height="{H}" font-family="Arial" font-size="12">',
           '<rect width="100%" height="100%" fill="white"/>',
           f'<text x="{W / 2}" y="28" text-anchor="middle" font-size="16" font-weight="bold" fill="#00338D">{html.escape(title)}</text>']
    if kind == "pie":
        vals = [float(v) for v in series[0]["values"]]
        total = sum(vals) or 1
        import math

        cx, cy, r, angle = W / 2 - 100, H / 2 + 15, 140, -math.pi / 2
        for i, (lab, v) in enumerate(zip(labels, vals)):
            a2 = angle + 2 * math.pi * v / total
            x1, y1 = cx + r * math.cos(angle), cy + r * math.sin(angle)
            x2, y2 = cx + r * math.cos(a2), cy + r * math.sin(a2)
            large = 1 if a2 - angle > math.pi else 0
            out.append(f'<path d="M{cx},{cy} L{x1:.1f},{y1:.1f} A{r},{r} 0 {large} 1 {x2:.1f},{y2:.1f} Z" fill="{PALETTE[i % 8]}" stroke="white" stroke-width="2"/>')
            out.append(f'<rect x="{W - 250}" y="{80 + i * 22}" width="12" height="12" fill="{PALETTE[i % 8]}"/>'
                       f'<text x="{W - 232}" y="{91 + i * 22}">{html.escape(str(lab))} ({v / total:.0%})</text>')
            angle = a2
        out.append("</svg>")
        return "".join(out)
    allv = [float(v) for sr in series for v in sr["values"]] or [0]
    vmax = max(allv + [0]) * 1.1 or 1
    vmin = min(allv + [0])
    rng = vmax - vmin or 1
    cw, chh = W - 2 * P, H - 2 * P

    def y(v):
        return P + chh - (float(v) - vmin) / rng * chh

    for t in range(5):
        val = vmin + rng * t / 4
        out.append(f'<line x1="{P}" x2="{W - P}" y1="{y(val):.1f}" y2="{y(val):.1f}" stroke="#e5e7eb"/>'
                   f'<text x="{P - 8}" y="{y(val) + 4:.1f}" text-anchor="end" fill="#555">{val:,.4g}</text>')
    nlab = max(1, len(labels))
    step = cw / nlab
    for i, lab in enumerate(labels):
        out.append(f'<text x="{P + step * i + step / 2:.1f}" y="{H - P + 18}" text-anchor="middle" fill="#333">{html.escape(str(lab))[:14]}</text>')
    for si, sr in enumerate(series):
        color = PALETTE[si % 8]
        if kind == "line":
            pts = " ".join(f"{P + step * i + step / 2:.1f},{y(v):.1f}" for i, v in enumerate(sr["values"]))
            out.append(f'<polyline points="{pts}" fill="none" stroke="{color}" stroke-width="2.5"/>')
            for i, v in enumerate(sr["values"]):
                out.append(f'<circle cx="{P + step * i + step / 2:.1f}" cy="{y(v):.1f}" r="3.5" fill="{color}"/>')
        else:
            bw = step * 0.8 / len(series)
            for i, v in enumerate(sr["values"]):
                x = P + step * i + step * 0.1 + bw * si
                out.append(f'<rect x="{x:.1f}" y="{min(y(v), y(0)):.1f}" width="{bw:.1f}" height="{abs(y(0) - y(v)):.1f}" fill="{color}" rx="2"/>')
        out.append(f'<rect x="{P + si * 140}" y="{H - 22}" width="12" height="12" fill="{color}"/>'
                   f'<text x="{P + si * 140 + 18}" y="{H - 12}">{html.escape(sr.get("name", f"Series {si + 1}"))}</text>')
    out.append("</svg>")
    return "".join(out)


@tool(
    "create_chart",
    "Create a bar, line or pie chart as an SVG image artifact (renders inline in chat).",
    obj(
        ["kind", "title", "labels", "series"],
        kind=s("bar | line | pie", enum=["bar", "line", "pie"]),
        title=s("Chart title"),
        labels=ARR_STR,
        series={"type": "array", "items": obj(["values"], name=s("Series name"), values={"type": "array", "items": {"type": "number"}})},
    ),
    risk="medium",
    category="Data",
)
async def create_chart(args, ctx):
    svg = svg_chart(args["kind"], args["title"], args["labels"], args["series"])
    return _saved(ctx, artifacts.save(f"{args['title']}.svg", svg, ctx.session_id, "create_chart"))


@tool(
    "generate_image",
    "Generate an image from a text prompt (Amazon Nova Canvas on Bedrock). Returns a PNG artifact.",
    obj(["prompt"], prompt=s("Detailed visual description"), width=n("default 1024"), height=n("default 1024")),
    risk="medium",
    category="Creative",
)
async def generate_image(args, ctx):
    from ..config import settings
    from ..llm import bedrock_image_sync

    if settings.get("model.provider") == "demo":
        raise ToolError("Image generation needs AWS Bedrock (Settings › Models).")
    try:
        png = await asyncio.to_thread(bedrock_image_sync, args["prompt"], int(args.get("width", 1024)), int(args.get("height", 1024)))
    except Exception as e:
        raise ToolError(f"Image generation failed: {e}")
    slug = re.sub(r"[^a-z0-9]+", "-", args["prompt"].lower())[:40].strip("-") or "image"
    return _saved(ctx, artifacts.save(f"{slug}.png", png, ctx.session_id, "generate_image"))


# ---------------------------------------------------------------------------
# Local files (gated by Settings > Directories)
# ---------------------------------------------------------------------------
@tool(
    "list_directory",
    "List files in an allowed folder (Settings › Directories, Mentor workspace, uploads, artifacts).",
    obj(["path"], path=s("Folder path"), pattern=s("Optional glob, e.g. '*.pdf'")),
    category="Files",
)
async def list_directory(args, ctx):
    try:
        p = resolve_allowed(args["path"])
    except PermissionError as e:
        raise ToolError(str(e))
    if not p.is_dir():
        raise ToolError(f"Not a folder: {p}")
    items = sorted(p.glob(args.get("pattern") or "*"))[:300]
    return "\n".join(f"{'[dir] ' if i.is_dir() else ''}{i.name}{'' if i.is_dir() else f'  ({i.stat().st_size:,} B)'}" for i in items) or "(empty)"


@tool(
    "read_document",
    "Read a local file's text (txt, md, csv, code, PDF, DOCX, XLSX, PPTX) from an allowed folder or an uploaded attachment.",
    obj(["path"], path=s("File path"), max_chars=n("Default 30000")),
    category="Files",
)
async def read_document(args, ctx):
    try:
        p = resolve_allowed(args["path"])
    except PermissionError as e:
        raise ToolError(str(e))
    if not p.is_file():
        raise ToolError(f"File not found: {p}")
    text = await asyncio.to_thread(ingest.extract_text, p)
    limit = int(args.get("max_chars", 30000))
    return text[:limit] + (f"\n\n... (truncated, {len(text):,} chars total)" if len(text) > limit else "")


@tool(
    "write_file",
    "Create or overwrite a text file in a folder with Read + write access. Relative paths go to the Mentor workspace.",
    obj(["path", "content"], path=s("File path"), content=s("Full content")),
    risk="high",
    category="Files",
)
async def write_file(args, ctx):
    from ..knowledge import ensure_workspace

    ensure_workspace()
    try:
        p = resolve_allowed(args["path"], write=True)
    except PermissionError as e:
        raise ToolError(str(e))
    p.parent.mkdir(parents=True, exist_ok=True)
    p.write_text(args["content"], "utf-8")
    return f"Wrote {len(args['content']):,} chars to {p}"


@tool(
    "search_files",
    "Search text inside files of an allowed folder (case-insensitive substring). Returns matching lines.",
    obj(["path", "query"], path=s("Folder path"), query=s("Text to find"), glob=s("Optional glob, default '**/*'")),
    category="Files",
)
async def search_files(args, ctx):
    try:
        root = resolve_allowed(args["path"])
    except PermissionError as e:
        raise ToolError(str(e))
    q = args["query"].lower()
    hits = []
    for p in root.glob(args.get("glob") or "**/*"):
        if not p.is_file() or ingest.skip_reason(p, root) or p.suffix.lower() not in ingest.TEXT_EXT:
            continue
        try:
            for ln, line in enumerate(p.read_text("utf-8", errors="ignore").splitlines(), 1):
                if q in line.lower():
                    hits.append(f"{p.relative_to(root)}:{ln}: {line.strip()[:200]}")
                    if len(hits) >= 100:
                        return "\n".join(hits) + "\n... (first 100 matches)"
        except OSError:
            continue
    return "\n".join(hits) or "No matches."


@tool(
    "analyze_table",
    "Profile a CSV or XLSX file: columns, types, row count, and summary stats for numeric columns.",
    obj(["path"], path=s("CSV/XLSX path"), sheet=s("Optional sheet name for XLSX")),
    category="Data",
)
async def analyze_table(args, ctx):
    try:
        p = resolve_allowed(args["path"])
    except PermissionError as e:
        raise ToolError(str(e))
    if p.suffix.lower() == ".xlsx":
        import openpyxl

        wb = openpyxl.load_workbook(str(p), read_only=True, data_only=True)
        ws = wb[args["sheet"]] if args.get("sheet") else wb.worksheets[0]
        rows = [list(r) for r in ws.iter_rows(values_only=True)]
    else:
        rows = list(csv.reader(io.StringIO(p.read_text("utf-8", errors="ignore"))))
    if not rows:
        return "Empty table."
    header = [str(h) for h in rows[0]]
    data = rows[1:]
    report = [f"Rows: {len(data):,}  Columns: {len(header)}"]
    for ci, name in enumerate(header):
        col = [r[ci] for r in data if ci < len(r) and r[ci] not in (None, "")]
        nums = []
        for v in col:
            try:
                nums.append(float(str(v).replace(",", "")))
            except ValueError:
                pass
        if col and len(nums) >= 0.8 * len(col):
            report.append(
                f"- {name} (numeric): n={len(nums)} min={min(nums):,.4g} max={max(nums):,.4g} "
                f"mean={statistics.fmean(nums):,.4g} median={statistics.median(nums):,.4g} sum={sum(nums):,.4g}"
            )
        else:
            uniq = len(set(map(str, col)))
            top = sorted({str(v) for v in col}, key=lambda v: -[str(x) for x in col].count(v))[:5]
            report.append(f"- {name} (text): non-empty={len(col)} unique={uniq} top={top}")
    return "\n".join(report)
