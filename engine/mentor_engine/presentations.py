"""Presentations: editable slide decks with real PowerPoint export/import."""

from __future__ import annotations

import io
from pathlib import Path
from typing import Any

from . import artifacts
from .db import new_id
from .state import db

LAYOUTS = ("title", "bullets", "two_column", "section", "quote", "closing")

THEMES = {
    "kpmg": {"bg": "FFFFFF", "title": "00338D", "text": "1E1E1E", "accent": "0091DA"},
    "midnight": {"bg": "0C1A3A", "title": "FFFFFF", "text": "DCE3F2", "accent": "00B8F5"},
    "minimal": {"bg": "FAFAFA", "title": "111111", "text": "333333", "accent": "7213EA"},
}


def _norm_slide(s: dict[str, Any]) -> dict[str, Any]:
    layout = s.get("layout", "bullets")
    return {
        "id": s.get("id") or new_id("sld_"),
        "layout": layout if layout in LAYOUTS else "bullets",
        "title": s.get("title", ""),
        "subtitle": s.get("subtitle", ""),
        "bullets": [str(b) for b in s.get("bullets", [])][:12],
        "right": [str(b) for b in s.get("right", [])][:12],
        "notes": s.get("notes", ""),
    }


def create(title: str, slides: list[dict[str, Any]], theme: str = "kpmg", session_id: str | None = None) -> dict[str, Any]:
    return db.put(
        "decks",
        {"title": title, "theme": theme if theme in THEMES else "kpmg", "slides": [_norm_slide(s) for s in slides], "session_id": session_id},
    )


def update(deck_id: str, patch: dict[str, Any]) -> dict[str, Any] | None:
    if "slides" in patch:
        patch["slides"] = [_norm_slide(s) for s in patch["slides"]]
    return db.patch("decks", deck_id, patch)


def _rgb(hexstr: str):
    from pptx.dml.color import RGBColor

    return RGBColor.from_string(hexstr)


def to_pptx_bytes(deck: dict[str, Any]) -> bytes:
    from pptx import Presentation
    from pptx.util import Inches, Pt

    theme = THEMES.get(deck.get("theme", "kpmg"), THEMES["kpmg"])
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)

    def style(tf, size, color, bold=False):
        for p in tf.paragraphs:
            for r in p.runs:
                r.font.size = Pt(size)
                r.font.color.rgb = _rgb(color)
                r.font.bold = bold

    def background(slide):
        fill = slide.background.fill
        fill.solid()
        fill.fore_color.rgb = _rgb(theme["bg"])

    for s in deck["slides"]:
        layout = s["layout"]
        if layout in ("title", "section", "closing"):
            slide = prs.slides.add_slide(prs.slide_layouts[0])
            background(slide)
            slide.shapes.title.text = s["title"]
            style(slide.shapes.title.text_frame, 40, theme["title"], True)
            if len(slide.placeholders) > 1:
                slide.placeholders[1].text = s.get("subtitle") or "\n".join(s["bullets"])
                style(slide.placeholders[1].text_frame, 20, theme["text"])
        elif layout == "two_column":
            slide = prs.slides.add_slide(prs.slide_layouts[3])
            background(slide)
            slide.shapes.title.text = s["title"]
            style(slide.shapes.title.text_frame, 30, theme["title"], True)
            for ph, items in ((slide.placeholders[1], s["bullets"]), (slide.placeholders[2], s["right"])):
                ph.text_frame.text = "\n".join(items)
                style(ph.text_frame, 18, theme["text"])
        elif layout == "quote":
            slide = prs.slides.add_slide(prs.slide_layouts[6])
            background(slide)
            box = slide.shapes.add_textbox(Inches(1.2), Inches(2.3), Inches(10.9), Inches(2.5))
            box.text_frame.word_wrap = True
            box.text_frame.text = f"“{s['title']}”"
            style(box.text_frame, 32, theme["title"], True)
            if s.get("subtitle"):
                p = box.text_frame.add_paragraph()
                p.text = f"— {s['subtitle']}"
                style(box.text_frame, 20, theme["text"])
        else:
            slide = prs.slides.add_slide(prs.slide_layouts[1])
            background(slide)
            slide.shapes.title.text = s["title"]
            style(slide.shapes.title.text_frame, 30, theme["title"], True)
            body = slide.placeholders[1].text_frame
            body.text = "\n".join(s["bullets"])
            style(body, 20, theme["text"])
        if s.get("notes"):
            slide.notes_slide.notes_text_frame.text = s["notes"]
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


def export(deck_id: str) -> dict[str, Any]:
    deck = db.get("decks", deck_id)
    if not deck:
        raise KeyError(deck_id)
    data = to_pptx_bytes(deck)
    name = "".join(c for c in deck["title"] if c.isalnum() or c in " -_").strip() or "deck"
    return artifacts.save(f"{name}.pptx", data, deck.get("session_id"), tool="presentations")


def import_pptx(name: str, data: bytes) -> dict[str, Any]:
    from pptx import Presentation

    prs = Presentation(io.BytesIO(data))
    slides = []
    for i, slide in enumerate(prs.slides):
        title = slide.shapes.title.text if slide.shapes.title is not None else ""
        bullets = []
        for shape in slide.shapes:
            if shape.has_text_frame and shape != slide.shapes.title:
                bullets += [p.text for p in shape.text_frame.paragraphs if p.text.strip()]
        notes = slide.notes_slide.notes_text_frame.text if slide.has_notes_slide else ""
        slides.append({"layout": "title" if i == 0 else "bullets", "title": title, "bullets": bullets, "notes": notes})
    return create(Path(name).stem, slides)
